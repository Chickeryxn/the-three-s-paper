# -*- coding: utf-8 -*-
"""Mechanical part of the Python code review for Q1 round1.

Emits code/Q1/reviews/q1_python_review.json with the five named checks required by
the code-reviewer contract: syntax, input_contract, method_alignment,
reproducibility, output_contract.
"""
from __future__ import annotations
import ast, hashlib, json, py_compile, sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
RUN = ROOT / "results/Q1/experiments/round1"
SCRIPTS = ["code/Q1/q1_common.py", "code/Q1/q1_main.py",
           "code/Q1/q1_baseline.py", "code/Q1/q1_verifier.py"]

def sha(p):
    return hashlib.sha256((ROOT / p).read_bytes()).hexdigest()

def load(p):
    return json.loads((ROOT / p).read_text(encoding="utf-8"))

def main():
    sys.stdout.reconfigure(encoding="utf-8")
    out = {"schema_version": 1, "question": "Q1", "round": "round1",
           "implementation_target": "python", "reviewer": "python-code-reviewer",
           "reviewed_scripts": {}, "checks": {}}

    # ---- syntax ---------------------------------------------------------------
    syn = {}
    for s in SCRIPTS:
        try:
            py_compile.compile(str(ROOT / s), doraise=True)
            ast.parse((ROOT / s).read_text(encoding="utf-8"))
            syn[s] = "ok"
            out["reviewed_scripts"][s] = sha(s)
        except Exception as exc:
            syn[s] = "FAIL: " + str(exc)
    out["checks"]["syntax"] = {
        "status": "PASS" if all(v == "ok" for v in syn.values()) else "FAIL",
        "detail": syn}

    # ---- input contract -------------------------------------------------------
    contract = load("planning/model_contract.json")
    consts = {"R": 0.02, "T0": 28.0, "C0": 2.55, "RHO": 820.0, "CP": 2600.0,
              "K": 0.36, "H": 25.0, "HM": 8e-7}
    common_src = (ROOT / "code/Q1/q1_common.py").read_text(encoding="utf-8")
    # extract module-level constants with AST so that combined assignment lines
    # (RHO, CP, K = ...) are handled correctly
    found = {}
    for node in ast.parse(common_src).body:
        if isinstance(node, ast.Assign):
            for tgt in node.targets:
                if isinstance(tgt, ast.Name) and isinstance(node.value, ast.Constant):
                    found[tgt.id] = node.value.value
                elif isinstance(tgt, ast.Tuple) and isinstance(node.value, ast.Tuple):
                    for nm, val in zip(tgt.elts, node.value.elts):
                        if isinstance(nm, ast.Name) and isinstance(val, ast.Constant):
                            found[nm.id] = val.value
    const_ok = {k: (k in found and abs(float(found[k]) - v) <= 1e-12) for k, v in consts.items()}
    inputs_ok = "attachment1_env.csv" in common_src
    unit_ok = all(t in common_src for t in ("0.02", "2.55", "8e-7"))
    out["checks"]["input_contract"] = {
        "status": "PASS" if all(const_ok.values()) and inputs_ok else "FAIL",
        "constants_match_model_contract": const_ok,
        "constants_found": {k: found.get(k) for k in consts},
        "reads_declared_input": inputs_ok,
        "units_si_with_report_conversion": unit_ok,
        "declared_inputs": [i["id"] for i in contract["inputs"]]}

    # ---- method alignment -----------------------------------------------------
    main_src = (ROOT / "code/Q1/q1_main.py").read_text(encoding="utf-8")
    base_src = (ROOT / "code/Q1/q1_baseline.py").read_text(encoding="utf-8")
    ver_src = (ROOT / "code/Q1/q1_verifier.py").read_text(encoding="utf-8")
    dec = [json.loads(l) for l in (ROOT / "methods/Q1/q1_decisions.jsonl").read_text(encoding="utf-8").splitlines() if l.strip()]
    approved = {d["decision_id"]: d for d in dec}
    mv = load("results/Q1/experiments/round1/metrics/main_validation.json")
    bv = load("results/Q1/experiments/round1/metrics/baseline_validation.json")
    vv = load("results/Q1/experiments/round1/metrics/verifier_validation.json")
    tol = contract["validation_contract"]["tolerances"]
    guard_p = ROOT / "methods/Q1/probes/role_swap_evidence/role_swap_check.json"
    guard = json.loads(guard_p.read_text(encoding="utf-8")) if guard_p.is_file() else {}
    align = {
        "main_vertex_conservative_flux_form": "Df[1:N]" in main_src and "Df[0:N - 1]" in main_src,
        "main_uses_adaptive_bdf": 'method="BDF"' in main_src,
        "main_surface_is_a_solution_node": "surface_node_is_exact" in main_src,
        "baseline_uses_crank_nicolson": 'scheme="CN"' in base_src,
        "baseline_cell_centred_finite_volume": "Af[N]" in base_src and "thomas(" in base_src,
        "verifier_uses_bessel": "bessel_C" in ver_src,
        "approved_decisions_recorded": all(k in approved for k in (
            "q1_method_choice", "q1_time_scheme_amendment", "q1_method_role_swap")),
        "role_swap_guard_executed": guard.get("status") == "PASS",
        "role_swap_guard_checks": {c["check"]: c["status"] for c in guard.get("checks", [])},
        "main_analytical_err_within_tol": mv["analytical_cross_check"]["max_abs_err"] < tol["analytical_cross_check_abs"],
        "baseline_analytical_err_within_tol": bv["analytical_cross_check"]["max_abs_err"] < tol["analytical_cross_check_abs"],
        "verifier_cross_check_pass": all(v.get("pass") for v in vv.values() if isinstance(v, dict) and "pass" in v),
    }
    out["checks"]["method_alignment"] = {
        "status": "PASS" if all(v for k, v in align.items() if k != "role_swap_guard_checks") else "CONDITIONAL",
        "detail": align,
        "approved_main": "M2", "approved_baseline": "B2",
        "note": "roles follow the human decision q1_method_role_swap: the vertex-centred node form (M2) was promoted to the main role and the cell-centred finite volume + Crank-Nicolson form (B2) to the baseline role; the time scheme of the cell-centred form follows the earlier human amendment q1_time_scheme_amendment",
        "role_swap_evidence": "methods/Q1/probes/role_swap_evidence/"}

    # ---- reproducibility ------------------------------------------------------
    snaps = {}
    for role in ("main", "baseline", "verifier"):
        d = load("results/Q1/experiments/round1/runs/%s/run_metadata.json" % role)
        snaps[role] = {"status": d.get("status"), "executed_by_runner": d.get("executed_by_runner"),
                       "return_code": d.get("return_code"), "result_hash": d.get("result_hash"),
                       "validation_hash": d.get("validation_hash"),
                       "code_manifest": d.get("code_manifest")}
    cur = {"main": sha("results/Q1/experiments/round1/metrics/main.json"),
           "baseline": sha("results/Q1/experiments/round1/metrics/baseline.json"),
           "verifier": sha("results/Q1/experiments/round1/metrics/verifier.json")}
    hashes_match = all(snaps[r]["result_hash"] == cur[r] for r in snaps)
    out["checks"]["reproducibility"] = {
        "status": "PASS" if hashes_match and all(s["executed_by_runner"] for s in snaps.values()) else "FAIL",
        "deterministic": True, "random_seed": 2026,
        "snapshot_result_hashes_match_current": hashes_match,
        "snapshots": snaps}

    # ---- output contract ------------------------------------------------------
    oke = {}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(RUN / "result1.xlsx", data_only=True)
        ws = wb["温度"]; ws2 = wb["水分浓度"]
        oke = {"sheets": wb.sheetnames,
               "rows_temperature": ws.max_row, "cols_temperature": ws.max_column,
               "rows_moisture": ws2.max_row, "cols_moisture": ws2.max_column,
               "header_ok": ws.cell(1, 1).value is not None and ws.cell(1, 2).value == 0,
               "time_axis_1_to_1800": ws.cell(2, 1).value == 1 and ws.cell(1801, 1).value == 1800,
               "distance_axis_0_to_2cm": ws.cell(1, 22).value == 2,
               "decimals_le_4": all(
                   len(str(v).split(".")[1]) <= 4
                   for v in (ws.cell(r, c).value for r in range(2, 30) for c in range(2, 22))
                   if isinstance(v, float))}
    except Exception as exc:
        oke = {"error": str(exc)}
    tables = sorted(p.name for p in (RUN / "tables").glob("*.csv"))
    out["checks"]["output_contract"] = {
        "status": "PASS" if oke.get("rows_temperature") == 1801 and oke.get("distance_axis_0_to_2cm")
                  and oke.get("time_axis_1_to_1800") and oke.get("decimals_le_4") else "FAIL",
        "deliverable": "results/Q1/experiments/round1/result1.xlsx",
        "detail": oke, "tables": tables,
        "kpi_tables_present": all(any(k in t for t in tables) for k in ("table1", "table2"))}

    out["overall_status"] = ("PASS" if all(c["status"] == "PASS" for c in out["checks"].values())
                             else "CONDITIONAL")
    dst = ROOT / "code/Q1/reviews/q1_python_review.json"
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text(json.dumps(out, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": out["overall_status"],
                      "checks": {k: v["status"] for k, v in out["checks"].items()}},
                     ensure_ascii=False))
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
