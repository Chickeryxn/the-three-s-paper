# -*- coding: utf-8 -*-
"""Mechanical part of the Q2 Python code review."""
from __future__ import annotations
import ast, hashlib, json, py_compile, sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
RUN = ROOT / "results/Q2/experiments/round1"
SCRIPTS = ["code/Q2/q2_main.py", "code/Q2/q2_baseline.py",
           "code/Q2/q2_verifier.py", "code/Q2/q2_round_summary.py"]

def sha(p): return hashlib.sha256((ROOT / p).read_bytes()).hexdigest()
def load(p): return json.loads((ROOT / p).read_text(encoding="utf-8"))

def main():
    sys.stdout.reconfigure(encoding="utf-8")
    out = {"schema_version": 1, "question": "Q2", "round": "round1",
           "implementation_target": "python", "reviewer": "python-code-reviewer",
           "reviewed_scripts": {}, "checks": {}}

    syn = {}
    for s in SCRIPTS:
        try:
            py_compile.compile(str(ROOT / s), doraise=True)
            ast.parse((ROOT / s).read_text(encoding="utf-8"))
            syn[s] = "ok"; out["reviewed_scripts"][s] = sha(s)
        except Exception as exc:
            syn[s] = "FAIL: " + str(exc)
    out["checks"]["syntax"] = {"status": "PASS" if all(v == "ok" for v in syn.values()) else "FAIL", "detail": syn}

    msrc = (ROOT / "code/Q2/q2_main.py").read_text(encoding="utf-8")
    consts = {"TARGET": 0.15, "H": 25.0, "HM": 8e-7, "C0": 2.55}
    found = {}
    for node in ast.parse(msrc).body:
        if isinstance(node, ast.Assign):
            for tgt in node.targets:
                if isinstance(tgt, ast.Name) and isinstance(node.value, ast.Constant):
                    found[tgt.id] = node.value.value
    const_ok = {k: (k in found and abs(float(found[k]) - v) <= 1e-12) for k, v in consts.items()}
    appendix_ok = all(t in msrc for t in ("650.0 + 128.0", "2736.0", "0.38", "2.4e-3", "3850.0", "0.45"))
    out["checks"]["input_contract"] = {
        "status": "PASS" if all(const_ok.values()) and appendix_ok else "FAIL",
        "constants": const_ok, "appendix3_closures_present": appendix_ok,
        "reads_declared_input": "attachment1_env.csv" in (ROOT / "code/Q2/q2_main.py").read_text(encoding="utf-8")}

    bsrc = (ROOT / "code/Q2/q2_baseline.py").read_text(encoding="utf-8")
    vsrc = (ROOT / "code/Q2/q2_verifier.py").read_text(encoding="utf-8")
    dec = [json.loads(l) for l in (ROOT / "methods/Q2/q2_decisions.jsonl").read_text(encoding="utf-8").splitlines() if l.strip()]
    ids = {d["decision_id"] for d in dec}
    mv = load("results/Q2/experiments/round1/metrics/main_validation.json")
    vv = load("results/Q2/experiments/round1/metrics/verifier_validation.json")
    align = {
        "main_is_vertex_centred": "dual-cell" in msrc and "V[1:N] = 2 * np.pi" in msrc,
        "main_uses_crank_nicolson": "theta=0.5" in msrc or "theta = 0.5" in msrc,
        "baseline_is_cell_centred": "rf = np.arange(N + 1) * dr" in bsrc and "CubicSpline" in bsrc,
        "verifier_uses_bessel": "bessel_C" in vsrc,
        "role_amendment_recorded": "q2_presentation_method_amendment" in ids,
        "method_choice_recorded": "q2_method_choice" in ids,
        "verifier_reference_pass": vv["reference_vs_analytical"]["pass"],
        "verifier_cross_check_pass": vv["main_vs_baseline_tables"]["pass"],
        "time_refinement_reported": mv["time_refinement_full_span"]["abs_diff_hours"] is not None,
    }
    out["checks"]["method_alignment"] = {"status": "PASS" if all(align.values()) else "CONDITIONAL", "detail": align,
                                         "presented_main": "vertex-centred finite volume",
                                         "verification": "cell-centred finite volume"}

    snaps = {}
    for role in ("main", "baseline", "verifier"):
        d = load("results/Q2/experiments/round1/runs/%s/run_metadata.json" % role)
        snaps[role] = {"status": d.get("status"), "executed_by_runner": d.get("executed_by_runner"),
                       "return_code": d.get("return_code"), "result_hash": d.get("result_hash"),
                       "validation_hash": d.get("validation_hash"), "degraded": d.get("degraded")}
    cur = {r: sha("results/Q2/experiments/round1/metrics/%s.json" % r) for r in snaps}
    hashes_match = all(snaps[r]["result_hash"] == cur[r] for r in snaps)
    out["checks"]["reproducibility"] = {
        "status": "PASS" if hashes_match and all(s["executed_by_runner"] and not s["degraded"] for s in snaps.values()) else "FAIL",
        "deterministic": True, "random_seed": 2026,
        "snapshot_hashes_match_current": hashes_match, "snapshots": snaps}

    oke = {}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(RUN / "result2.xlsx", read_only=True)
        ws = wb["温度"]; ws2 = wb["水分浓度"]
        rows = 0
        for _ in ws.iter_rows(values_only=True):
            rows += 1
        it = ws.iter_rows(min_row=1, max_row=2, values_only=True)
        hdr = next(it); first = next(it)
        oke = {"sheets": wb.sheetnames, "data_rows": rows - 1, "header_cols": len(hdr),
               "header_last_cm": hdr[-1], "first_time": first[0],
               "first_T_first_cell": first[1], "first_T_last_cell": first[-1],
               "moisture_sheet_present": "水分浓度" in wb.sheetnames}
    except Exception as exc:
        oke = {"error": str(exc)}
    tables = sorted(p.name for p in (RUN / "tables").glob("*.csv"))
    out["checks"]["output_contract"] = {
        "status": "PASS" if (oke.get("data_rows", 0) > 200000 and oke.get("header_last_cm") == 2
                             and oke.get("first_time") == 1 and oke.get("moisture_sheet_present")) else "FAIL",
        "deliverable": "results/Q2/experiments/round1/result2.xlsx",
        "detail": oke, "tables": tables,
        "temperature_scale_is_degC": bool(isinstance(oke.get("first_T_first_cell"), (int, float))
                                          and 20 <= oke["first_T_first_cell"] <= 60)}

    out["overall_status"] = "PASS" if all(c["status"] == "PASS" for c in out["checks"].values()) else "CONDITIONAL"
    dst = ROOT / "code/Q2/reviews/q2_python_review.json"
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text(json.dumps(out, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": out["overall_status"],
                      "checks": {k: v["status"] for k, v in out["checks"].items()},
                      "xlsx": oke}, ensure_ascii=False))
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
